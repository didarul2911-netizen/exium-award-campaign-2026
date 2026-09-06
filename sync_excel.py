"""
EXIUM MUPS - MASTER EXCEL SYNC SCRIPT FOR AWARD CHOICES
Run this script to update 'Exium_Award_Choice_Master_2026.xlsx' from a local backup JSON or downloaded export.
"""
import openpyxl
import json
import os
import sys

EXCEL_FILE = r"G:\Exium\2026\Award\May-Jun\Exium_Award_Choice_Master_2026.xlsx"
CHOICES_FILE = r"G:\Exium\2026\Award\May-Jun\saved_award_choices.json"

def sync_master_excel(choices_data):
    if not os.path.exists(EXCEL_FILE):
        print(f"Error: {EXCEL_FILE} not found. Please run generate_master_award_excel.py first.")
        return

    wb = openpyxl.load_workbook(EXCEL_FILE)
    
    # 1. Update May (Area Code Col 7, MIO Code Col 9, Voucher Col 15, Timestamp Col 16, Status Col 17)
    if "Award_May_2026" in wb.sheetnames:
        ws_may = wb["Award_May_2026"]
        updated_may = 0
        for r in range(2, ws_may.max_row + 1):
            area_code = str(ws_may.cell(r, 7).value or "").strip()
            mio_code = str(ws_may.cell(r, 9).value or "").strip()
            
            match_item = None
            for k, v in choices_data.items():
                if (k.startswith("MAY_") or v.get("month") == "May_2026") and (f"_{area_code}_" in k or v.get("area_code") == area_code):
                    if not mio_code or mio_code == "0" or mio_code in k or v.get("mio_code") == mio_code:
                        match_item = v
                        break
            
            if match_item:
                if match_item.get("voucher"):
                    ws_may.cell(r, 15, value=match_item["voucher"])
                    ws_may.cell(r, 16, value=match_item.get("timestamp", ""))
                    ws_may.cell(r, 17, value="Complete")
                    updated_may += 1
                elif match_item.get("timestamp"):
                    ws_may.cell(r, 15, value="")
                    ws_may.cell(r, 16, value=match_item.get("timestamp", ""))
                    ws_may.cell(r, 17, value="Pending")
                    updated_may += 1
        print(f"Updated {updated_may} records in Award_May_2026.")

    # 2. Update June (Area Code Col 7, MIO Code Col 9, Voucher Col 17, Timestamp Col 18, Status Col 19)
    if "Award_June_2026" in wb.sheetnames:
        ws_jun = wb["Award_June_2026"]
        updated_jun = 0
        for r in range(2, ws_jun.max_row + 1):
            area_code = str(ws_jun.cell(r, 7).value or "").strip()
            mio_code = str(ws_jun.cell(r, 9).value or "").strip()
            
            match_item = None
            for k, v in choices_data.items():
                if (k.startswith("JUN_") or v.get("month") == "June_2026") and (f"_{area_code}_" in k or v.get("area_code") == area_code):
                    if not mio_code or mio_code == "0" or mio_code in k or v.get("mio_code") == mio_code:
                        match_item = v
                        break
            
            if match_item:
                if match_item.get("voucher"):
                    ws_jun.cell(r, 17, value=match_item["voucher"])
                    ws_jun.cell(r, 18, value=match_item.get("timestamp", ""))
                    ws_jun.cell(r, 19, value="Complete")
                    updated_jun += 1
                elif match_item.get("timestamp"):
                    ws_jun.cell(r, 17, value="")
                    ws_jun.cell(r, 18, value=match_item.get("timestamp", ""))
                    ws_jun.cell(r, 19, value="Pending")
                    updated_jun += 1
        print(f"Updated {updated_jun} records in Award_June_2026.")

    wb.save(EXCEL_FILE)
    print(f"Master Excel '{EXCEL_FILE}' successfully synchronized and saved!")

if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else CHOICES_FILE
    if os.path.exists(src):
        with open(src, "r", encoding="utf-8") as f:
            data = json.load(f)
        sync_master_excel(data)
    else:
        print(f"Choices file '{src}' not found. Provide a JSON file path to sync.")
