import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_master_excel():
    folder = r"G:\Exium\2026\Award\May-Jun"
    output_path = os.path.join(folder, "Exium_Award_Choice_Master_2026.xlsx")
    
    # 1. Load FF list
    wb_ff = openpyxl.load_workbook(os.path.join(folder, "FF list.xlsx"), data_only=True)
    ws_ff = wb_ff["FF-RPL"]
    ff_by_mio = {}
    ff_by_terr = {}

    for r in range(2, ws_ff.max_row + 1):
        t_code = str(ws_ff.cell(r, 1).value or "").strip()
        t_name = str(ws_ff.cell(r, 2).value or "").strip()
        m_code = str(ws_ff.cell(r, 3).value or "").strip()
        m_name = str(ws_ff.cell(r, 4).value or "").strip()
        desig = str(ws_ff.cell(r, 5).value or "").strip()
        reg_code = str(ws_ff.cell(r, 6).value or "").strip()
        reg_name = str(ws_ff.cell(r, 7).value or "").strip()
        rh = str(ws_ff.cell(r, 8).value or "").strip()
        zone_code = str(ws_ff.cell(r, 9).value or "").strip()
        zone_name = str(ws_ff.cell(r, 10).value or "").strip()
        zh = str(ws_ff.cell(r, 11).value or "").strip()

        info = {
            "current_terr_code": t_code,
            "current_terr_name": t_name,
            "sap_mio_code": m_code,
            "current_mio_name": m_name,
            "desig": desig,
            "current_sap_reg_code": reg_code,
            "current_region": reg_name,
            "current_rh": rh,
            "current_sap_zone_code": zone_code,
            "current_zone": zone_name,
            "current_zh": zh
        }
        if m_code and m_code not in ["None", "0", ""]:
            ff_by_mio[m_code] = info
        if t_code:
            ff_by_terr[t_code] = info

    wb_master = openpyxl.Workbook()
    
    # Styles
    navy_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    gold_fill = PatternFill(start_color="D97706", end_color="D97706", fill_type="solid")
    emerald_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
    choice_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    transfer_fill = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")
    
    white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    regular_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )
    
    # ----------------------------------------------------
    # Sheet 1: Award_May_2026
    # ----------------------------------------------------
    ws_may = wb_master.create_sheet(title="Award_May_2026")
    wb_may_src = openpyxl.load_workbook(os.path.join(folder, "Exium Award Achiever List_May 2026.xlsx"), data_only=True)
    ws_may_src = wb_may_src["Award_May"]
    
    may_headers = [
        "SL", "Current Zone", "Current SAP Zone", "Current Region", "Current SAP Region", 
        "Current Regional Head", "Current Area Code", "Current Area Name", "SAP MIO Code", 
        "Sr./ MIO Name", "Designation", "No. of Rx", "Ach%", "Award Amount (BDT)", 
        "Selected Voucher", "Submission Timestamp", "Status",
        "Is Transferred", "Previous Area Code", "Previous Area Name", "Previous Region", "Previous Regional Head", "Previous Zone"
    ]
    
    ws_may.append(may_headers)
    for col_idx in range(1, len(may_headers) + 1):
        cell = ws_may.cell(1, col_idx)
        cell.font = white_bold
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if col_idx in [15, 16, 17]:
            cell.fill = gold_fill
        elif col_idx >= 18:
            cell.fill = PatternFill(start_color="9A3412", end_color="9A3412", fill_type="solid") # rust red
        else:
            cell.fill = navy_fill
    ws_may.row_dimensions[1].height = 28
    
    sl = 1
    for r in range(2, ws_may_src.max_row + 1):
        orig_zone = str(ws_may_src.cell(r, 1).value or "").strip()
        orig_reg = str(ws_may_src.cell(r, 2).value or "").strip()
        orig_rh = str(ws_may_src.cell(r, 3).value or "").strip()
        orig_ac = str(ws_may_src.cell(r, 4).value or "").strip()
        orig_an = str(ws_may_src.cell(r, 5).value or "").strip()
        orig_mc = str(ws_may_src.cell(r, 6).value or "").strip()
        orig_mn = str(ws_may_src.cell(r, 7).value or "").strip()
        rx = ws_may_src.cell(r, 8).value
        ach_pct = ws_may_src.cell(r, 9).value
        award = ws_may_src.cell(r, 10).value
        
        if not orig_ac and not orig_mn: continue

        ff = ff_by_mio.get(orig_mc) or ff_by_terr.get(orig_ac)
        is_tr = "NO"
        p_ac = p_an = p_reg = p_rh = p_z = ""

        if ff:
            c_z = ff["current_zone"]
            c_sz = ff["current_sap_zone_code"]
            c_reg = ff["current_region"]
            c_sr = ff["current_sap_reg_code"]
            c_rh = ff["current_rh"]
            c_ac = ff["current_terr_code"]
            c_an = ff["current_terr_name"]
            desig = ff["desig"]
            if c_ac != orig_ac or c_reg != orig_reg:
                is_tr = "YES"
                p_ac, p_an, p_reg, p_rh, p_z = orig_ac, orig_an, orig_reg, orig_rh, orig_zone
        else:
            c_z, c_sz, c_reg, c_sr, c_rh, c_ac, c_an = orig_zone, "", orig_reg, "", orig_rh, orig_ac, orig_an
            desig = "MIO"

        row_data = [
            sl, c_z, c_sz, c_reg, c_sr, c_rh, c_ac, c_an, orig_mc, orig_mn, desig,
            rx, ach_pct, award, "", "", "Not Started",
            is_tr, p_ac, p_an, p_reg, p_rh, p_z
        ]
        ws_may.append(row_data)
        curr_row = ws_may.max_row
        for c_idx in range(1, len(row_data) + 1):
            c = ws_may.cell(curr_row, c_idx)
            c.font = regular_font
            c.border = thin_border
            if c_idx in [1, 3, 5, 7, 9, 12, 14, 17, 18, 19]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 13:
                c.alignment = Alignment(horizontal="right", vertical="center")
                if isinstance(c.value, (int, float)): c.number_format = '0.00%'
            elif c_idx == 14:
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.number_format = '#,##0'
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")
            
            if c_idx in [15, 16, 17]:
                c.fill = choice_fill
            elif is_tr == "YES" and c_idx >= 18:
                c.fill = transfer_fill
        sl += 1

    # ----------------------------------------------------
    # Sheet 2: Award_June_2026
    # ----------------------------------------------------
    ws_jun = wb_master.create_sheet(title="Award_June_2026")
    wb_jun_src = openpyxl.load_workbook(os.path.join(folder, "Exium Award Achiever List_June 2026.xlsx"), data_only=True)
    ws_jun_src = wb_jun_src["Award_Jun"]
    
    jun_headers = [
        "SL", "Current Zone", "Current SAP Zone", "Current Region", "Current SAP Region", 
        "Current Regional Head", "Current Area Code", "Current Area Name", "SAP MIO Code", 
        "Sr./ MIO Name", "Designation", "No. of Rx", "Sales Val (Apr'26)", "Sales Val (Jun'26)", 
        "Growth% (Jun over Apr)", "Award Amount (BDT)", "Selected Voucher", "Submission Timestamp", "Status",
        "Is Transferred", "Previous Area Code", "Previous Area Name", "Previous Region", "Previous Regional Head", "Previous Zone"
    ]
    
    ws_jun.append(jun_headers)
    for col_idx in range(1, len(jun_headers) + 1):
        cell = ws_jun.cell(1, col_idx)
        cell.font = white_bold
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if col_idx in [17, 18, 19]:
            cell.fill = gold_fill
        elif col_idx >= 20:
            cell.fill = PatternFill(start_color="9A3412", end_color="9A3412", fill_type="solid")
        else:
            cell.fill = navy_fill
    ws_jun.row_dimensions[1].height = 28
    
    sl = 1
    for r in range(2, ws_jun_src.max_row + 1):
        orig_zone = str(ws_jun_src.cell(r, 1).value or "").strip()
        orig_reg = str(ws_jun_src.cell(r, 2).value or "").strip()
        orig_rh = str(ws_jun_src.cell(r, 3).value or "").strip()
        orig_ac = str(ws_jun_src.cell(r, 4).value or "").strip()
        orig_an = str(ws_jun_src.cell(r, 5).value or "").strip()
        orig_mc = str(ws_jun_src.cell(r, 6).value or "").strip()
        orig_mn = str(ws_jun_src.cell(r, 7).value or "").strip()
        rx = ws_jun_src.cell(r, 8).value
        sales_apr = ws_jun_src.cell(r, 9).value
        sales_jun = ws_jun_src.cell(r, 10).value
        growth = ws_jun_src.cell(r, 11).value
        award = ws_jun_src.cell(r, 12).value
        
        if not orig_ac and not orig_mn: continue

        ff = ff_by_mio.get(orig_mc) or ff_by_terr.get(orig_ac)
        is_tr = "NO"
        p_ac = p_an = p_reg = p_rh = p_z = ""

        if ff:
            c_z = ff["current_zone"]
            c_sz = ff["current_sap_zone_code"]
            c_reg = ff["current_region"]
            c_sr = ff["current_sap_reg_code"]
            c_rh = ff["current_rh"]
            c_ac = ff["current_terr_code"]
            c_an = ff["current_terr_name"]
            desig = ff["desig"]
            if c_ac != orig_ac or c_reg != orig_reg:
                is_tr = "YES"
                p_ac, p_an, p_reg, p_rh, p_z = orig_ac, orig_an, orig_reg, orig_rh, orig_zone
        else:
            c_z, c_sz, c_reg, c_sr, c_rh, c_ac, c_an = orig_zone, "", orig_reg, "", orig_rh, orig_ac, orig_an
            desig = "MIO"

        row_data = [
            sl, c_z, c_sz, c_reg, c_sr, c_rh, c_ac, c_an, orig_mc, orig_mn, desig,
            rx, sales_apr, sales_jun, growth, award, "", "", "Not Started",
            is_tr, p_ac, p_an, p_reg, p_rh, p_z
        ]
        ws_jun.append(row_data)
        curr_row = ws_jun.max_row
        for c_idx in range(1, len(row_data) + 1):
            c = ws_jun.cell(curr_row, c_idx)
            c.font = regular_font
            c.border = thin_border
            if c_idx in [1, 3, 5, 7, 9, 12, 16, 19, 20, 21]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx in [13, 14]:
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.number_format = '#,##0.00'
            elif c_idx == 15:
                c.alignment = Alignment(horizontal="right", vertical="center")
                if isinstance(c.value, (int, float)): c.number_format = '0.00%'
            elif c_idx == 16:
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.number_format = '#,##0'
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")
                
            if c_idx in [17, 18, 19]:
                c.fill = choice_fill
            elif is_tr == "YES" and c_idx >= 20:
                c.fill = transfer_fill
        sl += 1

    # ----------------------------------------------------
    # Sheet 3: Summary_Dashboard
    # ----------------------------------------------------
    ws_dash = wb_master.create_sheet(title="Summary_Dashboard", index=0)
    ws_dash.views.sheetView[0].showGridLines = True
    
    ws_dash.merge_cells("A1:G2")
    t_cell = ws_dash["A1"]
    t_cell.value = "EXIUM MUPS - SR. / MIO AWARD CATALOGUE CHOICE SUMMARY"
    t_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    t_cell.fill = navy_fill
    t_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    ws_dash["A4"] = "Campaign Metric"
    ws_dash["B4"] = "May 2026"
    ws_dash["C4"] = "June 2026"
    ws_dash["D4"] = "Total Combined"
    for col in ["A", "B", "C", "D"]:
        c = ws_dash[f"{col}4"]
        c.font = white_bold
        c.fill = navy_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        
    metrics = [
        ("Total Eligible Achievers", f"=COUNTA(Award_May_2026!A2:A{ws_may.max_row})", f"=COUNTA(Award_June_2026!A2:A{ws_jun.max_row})", "=B5+C5"),
        ("Completed Selections", f'=COUNTIF(Award_May_2026!Q2:Q{ws_may.max_row}, "Complete")', f'=COUNTIF(Award_June_2026!S2:S{ws_jun.max_row}, "Complete")', "=B6+C6"),
        ("Pending Selections", f'=COUNTIF(Award_May_2026!Q2:Q{ws_may.max_row}, "Not Started")', f'=COUNTIF(Award_June_2026!S2:S{ws_jun.max_row}, "Not Started")', "=B7+C7"),
        ("Completion Rate (%)", "=IF(B5>0, B6/B5, 0)", "=IF(C5>0, C6/C5, 0)", "=IF(D5>0, D6/D5, 0)")
    ]
    
    for idx, (label, f_may, f_jun, f_tot) in enumerate(metrics, start=5):
        ws_dash[f"A{idx}"] = label
        ws_dash[f"B{idx}"] = f_may
        ws_dash[f"C{idx}"] = f_jun
        ws_dash[f"D{idx}"] = f_tot
        for col in ["A", "B", "C", "D"]:
            c = ws_dash[f"{col}{idx}"]
            c.font = bold_font if col == "A" or idx == 8 else regular_font
            c.border = thin_border
            c.alignment = Alignment(horizontal="left" if col == "A" else "center", vertical="center")
            if idx == 8 and col in ["B", "C", "D"]:
                c.number_format = '0.0%'
    
    # Voucher Breakdown Section
    ws_dash["A11"] = "Voucher Brand"
    ws_dash["B11"] = "May Count"
    ws_dash["C11"] = "June Count"
    ws_dash["D11"] = "Total Selected"
    ws_dash["E11"] = "% Share"
    for col in ["A", "B", "C", "D", "E"]:
        c = ws_dash[f"{col}11"]
        c.font = white_bold
        c.fill = emerald_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        
    brands = [
        "Infinity Gift Voucher",
        "Apex Gift Voucher",
        "Bata Gift Voucher",
        "Aarong Gift Voucher",
        "Best Buy Gift Voucher",
        "Cats Eye Gift Voucher"
    ]
    
    for idx, brand in enumerate(brands, start=12):
        ws_dash[f"A{idx}"] = brand
        ws_dash[f"B{idx}"] = f'=COUNTIF(Award_May_2026!O2:O{ws_may.max_row}, "*{brand}*")'
        ws_dash[f"C{idx}"] = f'=COUNTIF(Award_June_2026!Q2:Q{ws_jun.max_row}, "*{brand}*")'
        ws_dash[f"D{idx}"] = f"=B{idx}+C{idx}"
        ws_dash[f"E{idx}"] = f"=IF($D$6>0, D{idx}/$D$6, 0)"
        for col in ["A", "B", "C", "D", "E"]:
            c = ws_dash[f"{col}{idx}"]
            c.border = thin_border
            c.alignment = Alignment(horizontal="left" if col == "A" else "center", vertical="center")
            if col == "E":
                c.number_format = '0.0%'
                
    # Auto-fit column widths
    for sheet in [ws_dash, ws_may, ws_jun]:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len and not cell.coordinate in ["A1", "B1", "C1", "D1", "E1", "F1", "G1"]:
                    max_len = len(val_str)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
    if "Sheet" in wb_master.sheetnames:
        wb_master.remove(wb_master["Sheet"])
        
    wb_master.save(output_path)
    print(f"Master Excel successfully updated with FF list mapping: {output_path}")

if __name__ == "__main__":
    generate_master_excel()
